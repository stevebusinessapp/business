from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Q, Sum, Count
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.template.loader import render_to_string
from django.forms import modelformset_factory
from decimal import Decimal
from .models import Quotation, QuotationItem, QuotationTemplate
from .forms import QuotationForm, QuotationItemFormSet, QuotationFilterForm, QuotationTemplateForm
from apps.clients.models import Client
from apps.core.models import CompanyProfile, format_currency, number_to_words
from apps.core.utils import get_company_context
import openpyxl
from xhtml2pdf import pisa
from io import BytesIO
import os
import base64
import json
from django.utils import timezone


def get_filtered_quotations(request):
    """Get filtered quotations based on search parameters"""
    quotations = Quotation.objects.filter(user=request.user).select_related('client', 'template')
    filter_form = QuotationFilterForm(request.GET)
    
    if filter_form.is_valid():
        search = filter_form.cleaned_data.get('search')
        status = filter_form.cleaned_data.get('status')
        date_from = filter_form.cleaned_data.get('date_from')
        date_to = filter_form.cleaned_data.get('date_to')
        client = filter_form.cleaned_data.get('client')
        
        if search:
            quotations = quotations.filter(
                Q(quotation_number__icontains=search) |
                Q(client__name__icontains=search) |
                Q(client__email__icontains=search) |
                Q(client__phone__icontains=search) |
                Q(notes__icontains=search) |
                Q(terms__icontains=search)
            )
        
        if status:
            quotations = quotations.filter(status=status)
        
        if date_from:
            quotations = quotations.filter(quotation_date__gte=date_from)
        
        if date_to:
            quotations = quotations.filter(quotation_date__lte=date_to)
        
        if client:
            quotations = quotations.filter(client=client)
    
    return quotations.order_by('-created_at')


@login_required
def quotation_list(request):
    """Display paginated list of quotations with filters"""
    quotations = get_filtered_quotations(request)
    filter_form = QuotationFilterForm(request.GET)
    
    # Update client choices for current user
    try:
        company_profile = request.user.company_profile
        if company_profile:
            filter_form.fields['client'].queryset = Client.objects.filter(company=company_profile)
    except:
        filter_form.fields['client'].queryset = Client.objects.none()
    
    paginator = Paginator(quotations, 25)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    # Calculate totals
    totals = quotations.aggregate(
        total_count=Count('id'),
        total_amount=Sum('grand_total') or 0,
        draft_count=Count('id', filter=Q(status='draft')),
        sent_count=Count('id', filter=Q(status='sent')),
        accepted_count=Count('id', filter=Q(status='accepted')),
        declined_count=Count('id', filter=Q(status='declined')),
        expired_count=Count('id', filter=Q(status='expired')),
    )
    
    # Get current template
    current_template = QuotationTemplate.objects.filter(user=request.user, is_default=True).first()
    if not current_template:
        current_template = QuotationTemplate.objects.filter(user=request.user).first()
        if not current_template:
            current_template = QuotationTemplate.objects.create(
                user=request.user,
                name="Default Quotation Template",
                description="Automatically created default template",
                is_default=True,
                primary_color="#1976d2",
                secondary_color="#f8f9fa",
                show_company_logo=True,
                show_company_details=True,
                show_bank_details=False,
                document_title="QUOTATION"
            )
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'page_obj': page_obj,
        'filter_form': filter_form,
        'total_quotations': totals['total_count'],
        'total_amount': totals['total_amount'],
        'formatted_total_amount': totals['total_amount'],
        'draft_count': totals['draft_count'],
        'sent_count': totals['sent_count'],
        'accepted_count': totals['accepted_count'],
        'declined_count': totals['declined_count'],
        'expired_count': totals['expired_count'],
        'current_template': current_template,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/quotation_list.html', context)


def get_quotation_context(quotation, user):
    """Get context data for quotation rendering"""
    # Get company profile data for display
    company_profile = None
    company_logo = None
    company_signature = None
    default_bank_account = None
    
    try:
        from apps.core.models import BankAccount
        company_profile = CompanyProfile.objects.get(user=user)
        
        if company_profile.logo:
            company_logo = company_profile.logo.url
        if company_profile.signature:
            company_signature = company_profile.signature.url
            
        # Get default bank account
        default_bank_account = BankAccount.objects.filter(
            company=company_profile, 
            is_default=True
        ).first()
    except:
        pass
    
    # Format currency and numbers
    currency_symbol = getattr(company_profile, 'currency_symbol', '$') if company_profile else '$'
    currency_code = getattr(company_profile, 'currency_code', 'USD') if company_profile else 'USD'
    total_words = number_to_words(quotation.grand_total, currency_name=currency_code)
    formatted_total = format_currency(quotation.grand_total, currency_symbol)
    
    return {
        'quotation': quotation,
        'company_profile': company_profile,
        'company_logo': company_logo,
        'company_signature': company_signature,
        'default_bank_account': default_bank_account,
        'currency_symbol': currency_symbol,
        'total_words': total_words,
        'formatted_total': formatted_total,
    }


@login_required
def quotation_detail(request, pk):
    """Display detailed view of a quotation"""
    quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
    context = get_quotation_context(quotation, request.user)
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    context.update({
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    })
    
    return render(request, 'quotations/quotation_detail.html', context)


@login_required
def quotation_create(request):
    """Create a new quotation"""
    if request.method == 'POST':
        form = QuotationForm(request.POST, user=request.user)
        formset = QuotationItemFormSet(request.POST, instance=Quotation(user=request.user))
        
        if form.is_valid() and formset.is_valid():
            quotation = form.save(commit=False)
            quotation.user = request.user
            quotation.save()
            
            formset.instance = quotation
            formset.save()
            
            quotation.calculate_totals()
            messages.success(request, f'Quotation {quotation.quotation_number} created successfully!')
            return redirect('quotations:quotation_detail', pk=quotation.pk)
    else:
        form = QuotationForm(user=request.user)
        formset = QuotationItemFormSet(instance=Quotation(user=request.user))
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'form': form,
        'formset': formset,
        'create': True,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/quotation_form.html', context)


@login_required
def quotation_edit(request, pk):
    """Edit an existing quotation"""
    quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
    
    if request.method == 'POST':
        form = QuotationForm(request.POST, instance=quotation, user=request.user)
        formset = QuotationItemFormSet(request.POST, instance=quotation)
        
        if form.is_valid() and formset.is_valid():
            quotation = form.save()
            formset.save()
            quotation.calculate_totals()
            messages.success(request, f'Quotation {quotation.quotation_number} updated successfully!')
            return redirect('quotations:quotation_detail', pk=quotation.pk)
    else:
        form = QuotationForm(instance=quotation, user=request.user)
        formset = QuotationItemFormSet(instance=quotation)
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'form': form,
        'formset': formset,
        'quotation': quotation,
        'create': False,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/quotation_form.html', context)


@login_required
def quotation_delete(request, pk):
    """Delete a quotation"""
    quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
    
    if request.method == 'POST':
        quotation_number = quotation.quotation_number
        quotation.delete()
        messages.success(request, f'Quotation {quotation_number} deleted successfully!')
        return redirect('quotations:quotation_list')
    
    context = {'quotation': quotation}
    return render(request, 'quotations/quotation_delete.html', context)


@login_required
def quotation_delete_ajax(request, pk):
    """Delete a quotation via AJAX"""
    if request.method == 'POST':
        try:
            quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
            quotation_number = quotation.quotation_number
            quotation.delete()
            return JsonResponse({
                'success': True,
                'message': f'Quotation {quotation_number} deleted successfully!'
            })
        except Exception as e:
            return JsonResponse({
                'success': False,
                'message': f'Error deleting quotation: {str(e)}'
            })
    
    return JsonResponse({'success': False, 'message': 'Invalid request method'})


@login_required
def quotation_pdf(request, pk):
    """Generate PDF for a quotation"""
    quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
    context = get_quotation_context(quotation, request.user)
    
    # Get template
    template = quotation.template
    if not template:
        template = QuotationTemplate.objects.filter(user=request.user, is_default=True).first()
        if not template:
            template = QuotationTemplate.objects.create(
                user=request.user,
                name="Default Template",
                is_default=True,
                primary_color="#1976d2",
                secondary_color="#f8f9fa"
            )
    
    context['template'] = template
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    context.update({
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    })
    
    try:
        html_string = render_to_string('quotations/quotation_pdf.html', context)
        result = BytesIO()
        pdf = pisa.pisaDocument(BytesIO(html_string.encode("UTF-8")), result)
        
        if not pdf.err:
            response = HttpResponse(result.getvalue(), content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="quotation_{quotation.quotation_number}.pdf"'
            return response
        else:
            return HttpResponse("Error generating PDF", content_type="text/plain", status=500)
    except Exception as e:
        return HttpResponse(f"Error generating PDF: {str(e)}", content_type="text/plain", status=500)


@login_required
def quotation_print(request, pk):
    """Print-friendly view of a quotation"""
    quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
    context = get_quotation_context(quotation, request.user)
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    context.update({
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    })
    
    return render(request, 'quotations/quotation_print.html', context)


@login_required
def template_list(request):
    """List all quotation templates"""
    templates = QuotationTemplate.objects.filter(user=request.user).order_by('-is_default', 'name')
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'templates': templates,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/template_list.html', context)


@login_required
def template_create(request):
    """Create a new quotation template"""
    if request.method == 'POST':
        form = QuotationTemplateForm(request.POST)
        if form.is_valid():
            template = form.save(commit=False)
            template.user = request.user
            template.save()
            messages.success(request, f'Template "{template.name}" created successfully!')
            return redirect('quotations:template_list')
    else:
        form = QuotationTemplateForm()
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'form': form,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/template_form.html', context)


@login_required
def template_edit(request, pk):
    """Edit a quotation template"""
    template = get_object_or_404(QuotationTemplate, pk=pk, user=request.user)
    
    if request.method == 'POST':
        form = QuotationTemplateForm(request.POST, instance=template)
        if form.is_valid():
            form.save()
            messages.success(request, f'Template "{template.name}" updated successfully!')
            return redirect('quotations:template_list')
    else:
        form = QuotationTemplateForm(instance=template)
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'form': form,
        'template': template,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/template_form.html', context)


@login_required
def template_detail(request, pk):
    """View template details"""
    template = get_object_or_404(QuotationTemplate, pk=pk, user=request.user)
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'template': template,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/template_detail.html', context)


@login_required
def template_delete(request, pk):
    """Delete a quotation template"""
    template = get_object_or_404(QuotationTemplate, pk=pk, user=request.user)
    
    if request.method == 'POST':
        template_name = template.name
        template.delete()
        messages.success(request, f'Template "{template_name}" deleted successfully!')
        return redirect('quotations:template_list')
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'template': template,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/template_delete.html', context)


@login_required
def template_duplicate(request, pk):
    """Duplicate a quotation template"""
    template = get_object_or_404(QuotationTemplate, pk=pk, user=request.user)
    
    if request.method == 'POST':
        new_template = QuotationTemplate.objects.create(
            user=request.user,
            name=f"{template.name} (Copy)",
            description=template.description,
            primary_color=template.primary_color,
            secondary_color=template.secondary_color,
            show_company_logo=template.show_company_logo,
            show_company_details=template.show_company_details,
            show_bank_details=template.show_bank_details,
            document_title=template.document_title,
            footer_text=template.footer_text,
            is_default=False
        )
        messages.success(request, f'Template "{new_template.name}" created successfully!')
        return redirect('quotations:template_list')
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'template': template,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/template_duplicate.html', context)


@login_required
def template_apply_to_all(request, pk):
    """Apply template to all quotations"""
    template = get_object_or_404(QuotationTemplate, pk=pk, user=request.user)
    
    if request.method == 'POST':
        updated_count = Quotation.objects.filter(user=request.user).update(template=template)
        messages.success(request, f'Template "{template.name}" applied to {updated_count} quotations!')
        return redirect('quotations:template_list')
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'template': template,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/template_apply.html', context)


@login_required
def api_templates(request):
    """API endpoint to get templates for AJAX requests"""
    templates = QuotationTemplate.objects.filter(user=request.user).values('id', 'name', 'is_default')
    return JsonResponse({'templates': list(templates)})


@login_required
def quotation_update_status(request, pk):
    """Update quotation status"""
    quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
    
    if request.method == 'POST':
        new_status = request.POST.get('status')
        if new_status in dict(Quotation.STATUS_CHOICES):
            old_status = quotation.get_status_display()
            quotation.status = new_status
            quotation.save()
            
            messages.success(request, f'Quotation status updated from {old_status} to {quotation.get_status_display()}')
            return redirect('quotations:quotation_detail', pk=quotation.pk)
        else:
            messages.error(request, 'Invalid status provided')
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'quotation': quotation,
        'status_choices': Quotation.STATUS_CHOICES,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/quotation_update_status.html', context)


@login_required
def convert_to_invoice(request, pk):
    """Convert quotation to invoice"""
    quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
    
    if request.method == 'POST':
        try:
            from apps.invoices.models import Invoice, InvoiceItem
            
            # Create invoice
            invoice = Invoice.objects.create(
                user=request.user,
                client=quotation.client,
                invoice_date=timezone.now().date(),
                due_date=timezone.now().date() + timezone.timedelta(days=30),
                subtotal=quotation.subtotal,
                total_tax=quotation.total_tax,
                total_discount=quotation.total_discount,
                shipping_fee=quotation.shipping_fee,
                other_charges=quotation.other_charges,
                grand_total=quotation.grand_total,
                notes=f"Converted from quotation {quotation.quotation_number}",
                status='unpaid'
            )
            
            # Copy line items
            for item in quotation.items.all():
                InvoiceItem.objects.create(
                    invoice=invoice,
                    product_service=item.product_service,
                    description=item.description,
                    quantity=item.quantity,
                    unit_price=item.unit_price,
                    line_total=item.line_total
                )
            
            # Update quotation
            quotation.converted_invoice = invoice
            quotation.conversion_date = timezone.now()
            quotation.status = 'accepted'
            quotation.save()
            
            messages.success(request, f'Quotation converted to invoice {invoice.invoice_number} successfully!')
            return redirect('invoices:invoice_detail', pk=invoice.pk)
            
        except Exception as e:
            messages.error(request, f'Error converting quotation: {str(e)}')
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'quotation': quotation,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/convert_to_invoice.html', context)


@login_required
def quotation_export_excel(request, pk=None):
    """Export quotations to Excel"""
    if pk:
        quotations = Quotation.objects.filter(pk=pk, user=request.user)
    else:
        quotations = get_filtered_quotations(request)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotations"
    
    # Headers
    headers = [
        "Quotation #", "Date", "Client", "Status", "Subtotal", 
        "Tax", "Discount", "Shipping", "Other Charges", "Grand Total"
    ]
    ws.append(headers)
    
    # Data
    for quotation in quotations:
        ws.append([
            quotation.quotation_number,
            quotation.quotation_date.strftime("%Y-%m-%d"),
            quotation.client.name if quotation.client else "No Client",
            quotation.get_status_display(),
            float(quotation.subtotal),
            float(quotation.total_tax),
            float(quotation.total_discount),
            float(quotation.shipping_fee),
            float(quotation.other_charges),
            float(quotation.grand_total),
        ])
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=quotations.xlsx'
    wb.save(response)
    return response


@login_required
def quotation_stats(request):
    """Get quotation statistics for dashboard"""
    quotations = Quotation.objects.filter(user=request.user)
    
    stats = {
        'total_quotations': quotations.count(),
        'total_value': float(quotations.aggregate(Sum('grand_total'))['grand_total__sum'] or 0),
        'draft_count': quotations.filter(status='draft').count(),
        'sent_count': quotations.filter(status='sent').count(),
        'accepted_count': quotations.filter(status='accepted').count(),
        'declined_count': quotations.filter(status='declined').count(),
        'expired_count': quotations.filter(status='expired').count(),
        'conversion_rate': 0,
    }
    
    if stats['total_quotations'] > 0:
        stats['conversion_rate'] = (stats['accepted_count'] / stats['total_quotations']) * 100
    
    return JsonResponse(stats)


@login_required
def quotation_duplicate(request, pk):
    """Duplicate a quotation"""
    quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
    
    if request.method == 'POST':
        try:
            # Create new quotation
            new_quotation = Quotation.objects.create(
                user=request.user,
                template=quotation.template,
                valid_until=quotation.valid_until,
                client=quotation.client,
                subtotal=quotation.subtotal,
                total_tax=quotation.total_tax,
                total_discount=quotation.total_discount,
                shipping_fee=quotation.shipping_fee,
                other_charges=quotation.other_charges,
                grand_total=quotation.grand_total,
                terms=quotation.terms,
                notes=quotation.notes,
                status='draft'
            )
            
            # Copy line items
            for item in quotation.items.all():
                QuotationItem.objects.create(
                    quotation=new_quotation,
                    product_service=item.product_service,
                    description=item.description,
                    quantity=item.quantity,
                    unit_price=item.unit_price,
                    line_total=item.line_total
                )
            
            messages.success(request, f'Quotation duplicated successfully! New quotation: {new_quotation.quotation_number}')
            return redirect('quotations:quotation_detail', pk=new_quotation.pk)
            
        except Exception as e:
            messages.error(request, f'Error duplicating quotation: {str(e)}')
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'quotation': quotation,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/quotation_duplicate.html', context)


@login_required
def quotation_bulk_actions(request):
    """Handle bulk actions on quotations"""
    if request.method == 'POST':
        action = request.POST.get('action')
        quotation_ids = request.POST.getlist('quotation_ids[]')
        
        if not quotation_ids:
            return JsonResponse({'success': False, 'message': 'No quotations selected'})
        
        quotations = Quotation.objects.filter(pk__in=quotation_ids, user=request.user)
        
        if action == 'delete':
            count = quotations.count()
            quotations.delete()
            return JsonResponse({
                'success': True, 
                'message': f'{count} quotation(s) deleted successfully!'
            })
        
        elif action == 'update_status':
            new_status = request.POST.get('status')
            if new_status in dict(Quotation.STATUS_CHOICES):
                count = quotations.update(status=new_status)
                return JsonResponse({
                    'success': True,
                    'message': f'Status updated for {count} quotation(s)'
                })
            else:
                return JsonResponse({
                    'success': False,
                    'message': 'Invalid status provided'
                })
        
        elif action == 'export':
            # Handle export logic
            return JsonResponse({
                'success': True,
                'message': f'Exporting {quotations.count()} quotation(s)'
            })
    
    return JsonResponse({'success': False, 'message': 'Invalid request'})


@login_required
def quotation_preview(request, pk):
    """Preview quotation before sending"""
    quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
    context = get_quotation_context(quotation, request.user)
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    context.update({
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    })
    
    return render(request, 'quotations/quotation_preview.html', context)


@login_required
def quotation_send_email(request, pk):
    """Send quotation via email"""
    quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
    
    if request.method == 'POST':
        # TODO: Implement email sending logic
        messages.info(request, 'Email feature not implemented yet.')
        return redirect('quotations:quotation_detail', pk=quotation.pk)
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'quotation': quotation,
        'company_profile': company_context['company_profile'],
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    return render(request, 'quotations/quotation_send_email.html', context)


@login_required
def quotation_export_pdf(request):
    """Export quotations list to PDF"""
    quotations = get_filtered_quotations(request)
    
    # Get company info and currency
    try:
        company_profile = request.user.company_profile
        company_logo_base64 = None
        currency_symbol = '$'
        if company_profile:
            currency_symbol = getattr(company_profile, 'currency_symbol', '$')
            if company_profile.logo:
                logo_path = company_profile.logo.path
                if os.path.isfile(logo_path):
                    with open(logo_path, 'rb') as img_file:
                        company_logo_base64 = 'data:image/png;base64,' + base64.b64encode(img_file.read()).decode('utf-8')
    except:
        company_profile = None
        company_logo_base64 = None
        currency_symbol = '$'
    
    # Calculate totals
    total_amount = quotations.aggregate(total=Sum('grand_total'))['total'] or 0
    
    # Get company context for currency
    company_context = get_company_context(request.user)
    
    context = {
        'quotations': quotations,
        'total_amount': total_amount,
        'company_profile': company_profile,
        'company_logo_base64': company_logo_base64,
        'currency_symbol': currency_symbol,
        'current_date': timezone.now(),
        'user_currency_symbol': company_context['currency_symbol'],
        'user_currency_code': company_context['currency_code'],
    }
    
    html_string = render_to_string('quotations/quotation_list_pdf.html', context)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html_string.encode("UTF-8")), result)
    
    if not pdf.err:
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename=quotations_list.pdf'
        return response
    else:
        return HttpResponse("Error generating PDF", content_type="text/plain", status=500)


@login_required
def quotation_debug(request):
    """Debug view to test quotation creation"""
    if request.method == 'POST':
        try:
            # Create a test quotation
            quotation = Quotation.objects.create(
                user=request.user,
                quotation_number="TEST-001",
                client=None,
                subtotal=Decimal('1000.00'),
                grand_total=Decimal('1000.00'),
                status='draft'
            )
            
            return JsonResponse({
                'success': True,
                'message': f'Test quotation created: {quotation.quotation_number}',
                'quotation_id': quotation.pk
            })
        except Exception as e:
            return JsonResponse({
                'success': False,
                'message': f'Error creating test quotation: {str(e)}'
            })
    
    return render(request, 'quotations/quotation_debug.html')


@login_required
def test_quotation_creation(request):
    """Test quotation creation with minimal data"""
    if request.method == 'POST':
        try:
            # Get or create a test client
            client, created = Client.objects.get_or_create(
                name="Test Client",
                defaults={
                    'email': 'test@example.com',
                    'phone': '1234567890',
                    'company': request.user.company_profile
                }
            )
            
            # Create quotation
            quotation = Quotation.objects.create(
                user=request.user,
                client=client,
                subtotal=Decimal('500.00'),
                grand_total=Decimal('500.00'),
                status='draft'
            )
            
            # Add a line item
            QuotationItem.objects.create(
                quotation=quotation,
                product_service="Test Product",
                description="Test Description",
                quantity=1,
                unit_price=Decimal('500.00'),
                line_total=Decimal('500.00')
            )
            
            return JsonResponse({
                'success': True,
                'message': f'Test quotation created successfully: {quotation.quotation_number}',
                'quotation_id': quotation.pk
            })
        except Exception as e:
            return JsonResponse({
                'success': False,
                'message': f'Error: {str(e)}'
            })
    
    return render(request, 'quotations/test_quotation.html')


@login_required
def test_simple_quotation(request):
    """Create a simple test quotation"""
    try:
        # Create quotation without client
        quotation = Quotation.objects.create(
            user=request.user,
            subtotal=Decimal('100.00'),
            grand_total=Decimal('100.00'),
            status='draft'
        )
        
        # Add line item
        QuotationItem.objects.create(
            quotation=quotation,
            product_service="Simple Test Product",
            quantity=1,
            unit_price=Decimal('100.00'),
            line_total=Decimal('100.00')
        )
        
        return JsonResponse({
            'success': True,
            'message': f'Simple quotation created: {quotation.quotation_number}',
            'quotation_id': quotation.pk
        })
    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'Error: {str(e)}'
        })


@login_required
def debug_quotation(request, pk):
    """Debug view for a specific quotation"""
    try:
        quotation = get_object_or_404(Quotation, pk=pk, user=request.user)
        
        debug_info = {
            'quotation_id': quotation.pk,
            'quotation_number': quotation.quotation_number,
            'user_id': quotation.user.id,
            'client_id': quotation.client.id if quotation.client else None,
            'template_id': quotation.template.id if quotation.template else None,
            'status': quotation.status,
            'subtotal': float(quotation.subtotal),
            'grand_total': float(quotation.grand_total),
            'items_count': quotation.items.count(),
            'created_at': quotation.created_at.isoformat(),
            'updated_at': quotation.updated_at.isoformat(),
        }
        
        return JsonResponse({
            'success': True,
            'debug_info': debug_info
        })
    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'Error: {str(e)}'
        })


@login_required
def test_quotation_state(request):
    """Test quotation state and calculations"""
    try:
        # Create a test quotation
        quotation = Quotation.objects.create(
            user=request.user,
            subtotal=Decimal('1000.00'),
            total_tax=Decimal('100.00'),
            total_discount=Decimal('50.00'),
            shipping_fee=Decimal('25.00'),
            other_charges=Decimal('10.00'),
            grand_total=Decimal('1085.00'),
            status='draft'
        )
        
        # Test calculations
        calculated_total = (
            quotation.subtotal + 
            quotation.total_tax + 
            quotation.shipping_fee + 
            quotation.other_charges - 
            quotation.total_discount
        )
        
        return JsonResponse({
            'success': True,
            'quotation_id': quotation.pk,
            'quotation_number': quotation.quotation_number,
            'subtotal': float(quotation.subtotal),
            'tax': float(quotation.total_tax),
            'discount': float(quotation.total_discount),
            'shipping': float(quotation.shipping_fee),
            'other_charges': float(quotation.other_charges),
            'grand_total': float(quotation.grand_total),
            'calculated_total': float(calculated_total),
            'calculation_correct': quotation.grand_total == calculated_total
        })
    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'Error: {str(e)}'
        })
